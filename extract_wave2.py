from docx import Document
p=r'E:\mmcc\jianqi\无声无息-拓展附魔（Without a sound Expanded Enchantments）-详细数值.docx'
d=Document(p)
for i,p in enumerate(d.paragraphs):
 if 410<=i<=470: print(i,p.text.strip())
for i,t in enumerate(d.tables):
 if i>=63: print('TABLE',i, [' | '.join(c.text for c in row.cells) for row in t.rows])