from docx import Document
p=r'E:\mmcc\jianqi\无声无息-拓展附魔（Without a sound Expanded Enchantments）-详细数值.docx'
d=Document(p)
for i,p in enumerate(d.paragraphs):
 if 300<=i<=350: print(i,p.text.strip())
for i,t in enumerate(d.tables):
 if 60<=i<=64:
  print('TABLE',i)
  for r in t.rows: print(' | '.join(c.text for c in r.cells))