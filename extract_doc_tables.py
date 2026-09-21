from docx import Document
p=r'E:\mmcc\jianqi\无声无息-拓展附魔（Without a sound Expanded Enchantments）-详细数值.docx'
d=Document(p)
for i,p in enumerate(d.paragraphs):
 t=p.text.strip()
 if t: print(f'P{i}: {t}')
for ti,t in enumerate(d.tables):
 print(f'TABLE {ti}')
 for row in t.rows:
  print(' | '.join(c.text.replace('\n',' / ').strip() for c in row.cells))