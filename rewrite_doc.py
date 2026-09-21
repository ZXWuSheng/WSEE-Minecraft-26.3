from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

src='无声无息-拓展附魔（Without a sound Expanded Enchantments）-详细数值.docx'
doc=Document(src)
sec=doc.sections[0]
sec.top_margin=Inches(.72); sec.bottom_margin=Inches(.72); sec.left_margin=Inches(.82); sec.right_margin=Inches(.82)
styles=doc.styles
styles['Normal'].font.name='Microsoft YaHei'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); styles['Normal'].font.size=Pt(10)
for name,size,color in [('Title',24,'17365D'),('Heading 1',16,'17365D'),('Heading 2',12,'2F5597'),('Heading 3',11,'2F5597')]:
    s=styles[name]; s.font.name='Microsoft YaHei'; s._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color)
for p in doc.paragraphs:
    p.paragraph_format.space_after=Pt(5); p.paragraph_format.line_spacing=1.12
    if p.style.name=='Title': p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(10)
    if p.style.name.startswith('Heading'): p.paragraph_format.keep_with_next=True; p.paragraph_format.space_before=Pt(10)
    for r in p.runs:
        r.font.name='Microsoft YaHei'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')

for table in doc.tables:
    table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.autofit=False
    for ri,row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(2); p.paragraph_format.line_spacing=1.0
                for r in p.runs:
                    r.font.name='Microsoft YaHei'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); r.font.size=Pt(9)
            if ri==0:
                tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'17365D'); tcPr.append(shd)
                for p in cell.paragraphs:
                    for r in p.runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255)

doc.save(src)
print('rewritten')
